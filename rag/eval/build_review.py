"""Phase 2, step 3: the spreadsheet for reviewing the drafted questions.

Joins every draft to the chunk it was written from, marks the answer inside the
chunk as >>>answer<<< in bold red so it is found at a glance, and re-checks
that every marked answer really exists in its chunk.

The reviewer picks a verdict from a dropdown: ok, fix or drop. For fix, the
corrected question goes in fixed_question. Only draft_id, verdict,
fixed_question and notes are read back, so sorting or filtering the sheet is
safe, and a cell the spreadsheet reformats cannot leak into the gold set.

Written as .xlsx rather than .csv. 82 of the 140 chunks contain line breaks,
and spreadsheet programs disagree on how to read CSV: Excel takes the
delimiter from the Windows locale, and WPS opened this sheet as a single
column. An .xlsx stores cells rather than text, so nothing has to be guessed.
The CSV writer is kept, switched off in main().

A sheet that already holds verdicts is never overwritten, so re-running this
cannot erase a review.

Output: data/rag/gold_review.xlsx. Gitignored: the reviewed sheet becomes the
committed gold set in the next step.

Usage:
    python -m rag.eval.build_review
"""
from __future__ import annotations

import csv
import json
import os
from collections import Counter

import duckdb
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RAGDATA = "D:/capstone/data/rag"
DRAFTS = f"{RAGDATA}/gold_drafts.jsonl"
CONTENT = f"{RAGDATA}/gold_candidates.jsonl"
VERIFY = f"{RAGDATA}/gold_verify_candidates.jsonl"
NEGATIVE = f"{RAGDATA}/gold_negative_candidates.jsonl"
CHUNKS = f"{RAGDATA}/chunks.parquet"
OUT = f"{RAGDATA}/gold_review.xlsx"
OUT_CSV = f"{RAGDATA}/gold_review.csv"

# Reviewer columns first, then left to right what judging a question needs:
# the drug, the question, its answer, and the chunk the answer sits in.
COLUMNS = ["draft_id", "verdict", "fixed_question", "notes",
           "qtype", "drug", "question", "answer_span",
           "chunk_with_answer_marked", "rationale",
           "brand", "section", "negative_kind", "candidate_id"]
REVIEWER = ("verdict", "fixed_question", "notes")

# Column widths, in characters.
WIDTHS = {"draft_id": 9, "verdict": 9, "fixed_question": 30, "notes": 20,
          "qtype": 11, "drug": 16, "question": 40, "answer_span": 24,
          "chunk_with_answer_marked": 90, "rationale": 36,
          "brand": 14, "section": 20, "negative_kind": 18, "candidate_id": 12}

VERDICTS = ("ok", "fix", "drop")
NO_CHUNK = "(no chunk: the correct answer is 'outside indexed scope')"
ANSWER_FONT = InlineFont(b=True, color="C00000")


def excel_delimiter() -> str:
    """Excel splits a CSV on the Windows list separator, which is ';' on many
    European locales. Reading it means a double-click opens the file cleanly."""
    try:
        import winreg
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, r"Control Panel\International") as key:
            return winreg.QueryValueEx(key, "sList")[0] or ","
    except (ImportError, OSError):
        return ","


def safe(value: str) -> str:
    """Excel treats a cell starting with = + - @ as a formula, and label text
    can start with '- Bleeding', which would display as #NAME?."""
    return "'" + value if value[:1] in ("=", "+", "-", "@") else value


def load(path: str, key: str) -> dict:
    with open(path, encoding="utf-8") as fh:
        return {row[key]: row for row in map(json.loads, fh)}


def build_rows() -> tuple[list[dict], list[str]]:
    """One row per draft, joined to its chunk. Also returns the draft_ids whose
    answer was not found in their chunk."""
    with open(DRAFTS, encoding="utf-8") as fh:
        drafts = [json.loads(line) for line in fh]
    content = load(CONTENT, "candidate_id")
    verify = load(VERIFY, "vcand_id")
    negative = load(NEGATIVE, "ncand_id")

    ids = [v["chunk_id"] for v in verify.values()]
    text_by_chunk = dict(duckdb.connect().execute(
        f"SELECT chunk_id, text FROM '{CHUNKS}' WHERE list_contains(?, chunk_id)",
        [ids]).fetchall())

    rows, missing = [], []
    for d in drafts:
        cid = d.get("candidate_id") or ""
        if cid.startswith("c"):
            c = content[cid]
            text, drug, brand, section = c["text"], c["generic_name"], c["brand_name"], c["section"]
        elif cid.startswith("v"):
            v = verify[cid]
            text = text_by_chunk.get(v["chunk_id"], "")
            drug, brand, section = v["generic_name"], v["brand_name"], v["section"]
        else:
            text, brand, section = "", "", ""
            drug = negative[cid]["drug_name"] if cid.startswith("n") else d.get("drug", "")

        # parts = (before, answer, after) so the xlsx can colour the answer.
        # _text is the plain chunk, which finalize_gold checks corrected spans against.
        span = d.get("answer_span") or ""
        if not span:
            marked, parts = NO_CHUNK, None
        elif span in text:
            i = text.index(span)
            parts = (text[:i], span, text[i + len(span):])
            marked = f"{parts[0]}>>>{span}<<<{parts[2]}"
        else:
            marked, parts = text, None
            missing.append(d["draft_id"])

        rows.append({
            "draft_id": d["draft_id"], "verdict": "", "fixed_question": "", "notes": "",
            "qtype": d["qtype"], "question": d["question"], "answer_span": span,
            "drug": drug or "", "brand": brand or "", "section": section,
            "chunk_with_answer_marked": marked,
            "rationale": d.get("rationale", ""),
            "negative_kind": d.get("negative_kind") or "", "candidate_id": cid,
            "_parts": parts,
            "_text": text,
        })
    return rows, missing


def verdicts_in(path: str) -> int:
    """How many verdicts an existing review sheet already holds."""
    if not os.path.exists(path):
        return 0
    wb = load_workbook(path, read_only=True)
    try:
        cells = wb.active.iter_rows(values_only=True)
        header = list(next(cells, ()))
        if "verdict" not in header:
            return 0
        col = header.index("verdict")
        return sum(1 for r in cells if col < len(r) and str(r[col] or "").strip())
    finally:
        wb.close()


def xlsx_text(value: str) -> str:
    """XML cannot hold most control characters, and openpyxl refuses them."""
    return ILLEGAL_CHARACTERS_RE.sub("", value)


def write_xlsx(rows: list[dict], path: str) -> None:
    """The review sheet, ready to use: top row frozen, filter on, every cell
    wrapped, the answer in bold red, and a dropdown for the verdict."""
    wb = Workbook()
    ws = wb.active
    ws.title = "review"
    wrap = Alignment(wrap_text=True, vertical="top")
    reviewer_fill = PatternFill("solid", fgColor="FFF2CC")

    for j, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=j, value=name)
        cell.font = Font(bold=True)
        if name in REVIEWER:
            cell.fill = reviewer_fill
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS[name]

    for i, row in enumerate(rows, start=2):
        for j, name in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=i, column=j)
            parts = row["_parts"] if name == "chunk_with_answer_marked" else None
            if parts:
                before, span, after = (xlsx_text(p) for p in parts)
                runs = [before] if before else []
                runs.append(TextBlock(ANSWER_FONT, f">>>{span}<<<"))
                if after:
                    runs.append(after)
                cell.value = CellRichText(*runs)
            else:
                text = xlsx_text(row[name])
                cell.value = text or None
                if text:
                    # openpyxl reads "=..." as a formula and "#N/A" as an
                    # error. Label text is always just text.
                    cell.data_type = "s"
            cell.alignment = wrap

    last = len(rows) + 1
    verdict_col = get_column_letter(COLUMNS.index("verdict") + 1)
    choices = DataValidation(type="list", formula1='"' + ",".join(VERDICTS) + '"',
                             allow_blank=True, showErrorMessage=True,
                             errorTitle="verdict", error="Pick ok, fix or drop.")
    ws.add_data_validation(choices)
    choices.add(f"{verdict_col}2:{verdict_col}{last}")
    ws.freeze_panes = "C2"  # header row, draft_id and verdict stay in view
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"

    try:
        wb.save(path)
    except PermissionError:
        raise SystemExit(f"cannot write {path}: it is open in WPS or Excel. "
                         "Close it and run this again.")


def write_csv(rows: list[dict], path: str) -> None:
    """The same sheet as CSV, split on the delimiter Excel expects here."""
    delim = excel_delimiter()
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=COLUMNS, delimiter=delim,
                                extrasaction="ignore")
        writer.writeheader()
        writer.writerows({k: safe(v) if isinstance(v, str) else v
                          for k, v in row.items()} for row in rows)
    print(f"CSV copy written to {path}  (delimiter {delim!r})")


def main() -> None:
    done = verdicts_in(OUT)
    if done:
        raise SystemExit(f"{OUT} already has a verdict on {done} of its rows, and "
                         "rebuilding it would erase the review. Rename or move it first.")

    rows, missing = build_rows()
    write_xlsx(rows, OUT)
    # CSV is the more common format, but it is switched off: line breaks inside
    # cells do not survive every spreadsheet program. Uncomment to write it too.
    # write_csv(rows, OUT_CSV)

    print(f"{len(rows)} questions written to {OUT}")
    for qtype, n in Counter(r["qtype"] for r in rows).most_common():
        print(f"  {qtype:<12} {n}")
    print(f"answers not found in their chunk: {len(missing)} {missing or ''}")


if __name__ == "__main__":
    main()
