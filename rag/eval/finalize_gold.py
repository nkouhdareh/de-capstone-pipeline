"""Phase 2, step 4: turn the reviewed sheet into the committed gold set.

Reads back only four columns: draft_id, verdict, fixed_question and notes. The
question, the answer span and the anchor all come from the draft and candidate
files, so nothing a spreadsheet reformats can reach the gold set, and sorting
or filtering the sheet is harmless.

Verdicts: ok keeps the drafted question, fix keeps the reviewer's wording, drop
throws the row away. A reviewer can also correct an answer span, by writing the
corrected text between >>> and <<< in notes, which is the same marker the sheet
uses. Every correction is checked against the chunk word for word, so a typo
fails the run instead of quietly poisoning the gold set.

Anchors are (set_id, section, answer_span), never chunk ids. Chunk ids change
whenever the chunker changes, and this file has to outlive that: re-chunk,
re-resolve, same gold set. Negatives carry no anchor, because their answer is
not in the corpus at all. No drug_key anywhere, per the Phase 1 constraint.

Split: 60% dev, 40% test, stratified by question type, so a type with 15
questions is not split 3/12 by luck. Tune on dev, report on test once. The
order inside each type comes from md5(draft_id + SALT), which is stable across
machines and Python versions, unlike a seeded shuffle.

Output: rag/eval/gold/gold.jsonl, the one file in this phase that is committed.

Usage:
    python -m rag.eval.finalize_gold
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from rag.eval.build_review import build_rows

RAGDATA = "D:/capstone/data/rag"
REVIEWED = f"{RAGDATA}/gold_reviewed.xlsx"   # a .csv export is read too
DRAFTS = f"{RAGDATA}/gold_drafts.jsonl"
CONTENT = f"{RAGDATA}/gold_candidates.jsonl"
VERIFY = f"{RAGDATA}/gold_verify_candidates.jsonl"
OUT = Path(__file__).resolve().parent / "gold" / "gold.jsonl"

SALT = "gold-v1"        # changing this reshuffles the dev/test split
DEV_SHARE = 0.6
KEEP = ("ok", "fix")
VERDICTS = ("ok", "fix", "drop")
READ_BACK = ("draft_id", "verdict", "fixed_question", "notes")
CORRECTED_SPAN = re.compile(r">>>(.+?)<<<", re.DOTALL)


def clean(value) -> str:
    return str(value).strip() if value is not None else ""


def load_jsonl(path: str, key: str) -> dict:
    with open(path, encoding="utf-8") as fh:
        return {row[key]: row for row in map(json.loads, fh)}


def read_xlsx(path: str) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb.active.iter_rows(values_only=True)
        header = [clean(h) for h in next(rows, ())]
        return [dict(zip(header, (clean(v) for v in row))) for row in rows]
    finally:
        wb.close()


def read_csv(path: str) -> list[dict]:
    """Excel and WPS each write CSV in whatever the machine's locale says, so
    the delimiter is read off the header line rather than assumed, and the
    encoding is tried twice: a UTF-8 export first, then a plain Windows one."""
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            with open(path, encoding=encoding, newline="") as fh:
                text = fh.read()
            break
        except UnicodeDecodeError:
            continue
    else:
        raise SystemExit(f"{path}: not readable as utf-8 or cp1252")
    header = text.splitlines()[0] if text else ""
    delimiter = max(",;\t|", key=header.count)
    reader = csv.DictReader(text.splitlines(True), delimiter=delimiter)
    return [{clean(k): clean(v) for k, v in row.items()} for row in reader]


def read_review(path: str) -> dict:
    """draft_id -> the four reviewer columns. Everything else is ignored."""
    rows = read_xlsx(path) if path.lower().endswith(".xlsx") else read_csv(path)
    review = {}
    for row in rows:
        draft_id = row.get("draft_id", "")
        if not draft_id:
            continue
        if draft_id in review:
            raise SystemExit(f"{path}: {draft_id} appears on two rows")
        review[draft_id] = {name: row.get(name, "") for name in READ_BACK}
    return review


def stratified_split(qids: list[str]) -> dict:
    """Deterministic 60/40 inside one question type."""
    order = sorted(qids, key=lambda q: hashlib.md5(f"{q}|{SALT}".encode()).hexdigest())
    n_dev = round(len(order) * DEV_SHARE)
    return {q: ("dev" if i < n_dev else "test") for i, q in enumerate(order)}


def build_gold(drafts, content, verify, chunks, review):
    """Returns the kept records and a list of problems. Any problem at all means
    nothing is written: a gold set is not worth having if part of it is guessed."""
    kept, problems = [], []
    for draft_id, draft in drafts.items():
        row = review[draft_id]
        verdict = row["verdict"].lower()
        if verdict not in VERDICTS:
            problems.append(f"{draft_id}: verdict {row['verdict']!r} is not ok, fix or drop")
            continue
        if verdict not in KEEP:
            continue

        question = row["fixed_question"] or draft["question"]
        correction = CORRECTED_SPAN.search(row["notes"])
        span = correction.group(1).strip() if correction else (draft.get("answer_span") or "")
        if verdict == "fix" and not row["fixed_question"] and not correction:
            problems.append(
                f"{draft_id}: marked fix, but neither the question nor the span was corrected")
            continue

        anchor = None
        if draft["qtype"] != "negative":
            candidate_id = draft.get("candidate_id") or ""
            candidate = content.get(candidate_id) or verify.get(candidate_id)
            if candidate is None:
                problems.append(f"{draft_id}: no candidate chunk to anchor on")
                continue
            if not span or span not in chunks[draft_id]["_text"]:
                problems.append(f"{draft_id}: the answer span is not in its chunk word for word")
                continue
            anchor = {"set_id": candidate["set_id"],
                      "section": candidate["section"],
                      "answer_span": span}

        kept.append({
            "qid": draft_id,
            "split": "",
            "qtype": draft["qtype"],
            "question": question,
            "anchor": anchor,
            "drug": chunks[draft_id]["drug"] or None,
            "brand": chunks[draft_id]["brand"] or None,
            "negative_kind": draft.get("negative_kind") or None,
            "provenance": {
                "source": draft.get("source", ""),
                "verdict": verdict,
                "candidate_id": draft.get("candidate_id") or None,
                "question_edited": bool(row["fixed_question"]),
                "span_edited": bool(correction),
            },
        })
    return kept, problems


def main() -> None:
    drafts = load_jsonl(DRAFTS, "draft_id")
    content = load_jsonl(CONTENT, "candidate_id")
    verify = load_jsonl(VERIFY, "vcand_id")
    chunks = {row["draft_id"]: row for row in build_rows()[0]}
    review = read_review(REVIEWED)

    absent = sorted(set(drafts) - set(review))
    unknown = sorted(set(review) - set(drafts))
    blank = sorted(d for d in review if not review[d]["verdict"])
    if absent or unknown or blank:
        raise SystemExit(f"{REVIEWED} does not match the drafts.\n"
                         f"  missing from the sheet: {absent or 'none'}\n"
                         f"  not a known draft_id:   {unknown or 'none'}\n"
                         f"  no verdict:             {blank or 'none'}")

    kept, problems = build_gold(drafts, content, verify, chunks, review)
    if problems:
        raise SystemExit("nothing was written, because:\n  " + "\n  ".join(problems))

    by_type = defaultdict(list)
    for record in kept:
        by_type[record["qtype"]].append(record["qid"])
    split = {}
    for qids in by_type.values():
        split.update(stratified_split(qids))
    for record in kept:
        record["split"] = split[record["qid"]]

    kept.sort(key=lambda r: r["qid"])
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with open(OUT, "w", encoding="utf-8", newline="\n") as fh:
        for record in kept:
            fh.write(json.dumps(record, ensure_ascii=False) + "\n")

    dropped = sorted(d for d in review if review[d]["verdict"].lower() == "drop")
    edits = Counter(k for r in kept for k in ("question_edited", "span_edited")
                    if r["provenance"][k])
    print(f"{len(kept)} questions written to {OUT}")
    print(f"  dropped by the reviewer:  {len(dropped)} {dropped or ''}")
    print(f"  questions rewritten:      {edits['question_edited']}")
    print(f"  answer spans corrected:   {edits['span_edited']}")
    print()
    counts = Counter((r["qtype"], r["split"]) for r in kept)
    print(f"  {'qtype':<12} {'dev':>4} {'test':>5} {'all':>5}  share")
    for qtype in sorted(by_type):
        dev, test = counts[(qtype, "dev")], counts[(qtype, "test")]
        print(f"  {qtype:<12} {dev:>4} {test:>5} {dev + test:>5}  {(dev + test) / len(kept):>5.0%}")
    dev = sum(1 for r in kept if r["split"] == "dev")
    print(f"  {'TOTAL':<12} {dev:>4} {len(kept) - dev:>5} {len(kept):>5}  "
          f"{dev / len(kept):.0%} dev / {(len(kept) - dev) / len(kept):.0%} test")


if __name__ == "__main__":
    main()
